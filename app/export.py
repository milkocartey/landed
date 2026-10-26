"""Export the tailored CV and cover letter as Word documents."""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, RGBColor

TITLES = {"experience": "Experience", "projects": "Projects", "education": "Education",
          "activities": "Activities", "languages": "Languages"}
ACCENT = RGBColor(0x2B, 0x2A, 0x6E)


def _base() -> Document:
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    return d


def letter_docx(text: str) -> bytes:
    d = _base()
    for para in text.split("\n\n"):
        d.add_paragraph(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def cv_docx(cv: dict) -> bytes:
    d = _base()
    header = cv.get("header") or []
    if header:
        h = d.add_paragraph()
        r = h.add_run(header[0])
        r.bold, r.font.size, r.font.color.rgb = True, Pt(18), ACCENT
        for line in header[1:]:
            d.add_paragraph(line)

    def heading(t):
        p = d.add_paragraph()
        r = p.add_run(t.upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(11), ACCENT

    heading("Profile")
    d.add_paragraph(cv.get("summary", ""))
    heading("Skills")
    d.add_paragraph(" · ".join(cv.get("skills", [])))
    for key, lines in cv.get("sections", {}).items():
        heading(TITLES.get(key, key.title()))
        for line in lines:
            d.add_paragraph(line.lstrip("-•*–▪ "), style="List Bullet" if line[:1] in "-•*–▪" else None)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
